# -*- coding: utf-8 -*-
"""Quality gate for the generated statement sets.

  1. Every output workbook opens with openpyxl and its cell styles are a subset
     of the source workbook's styles; five sampled cells are compared field by
     field between source and output.
  2. Piedmont is diffed line by line against the Boom fixture.
  3. No em dash in any produced text.
"""
import json, os, sys
import openpyxl

DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(DIR, '..', '..', 'preview-site', 'builds', 'statements-0915')
SRC = os.path.join(DIR, 'template-workbook.xlsx')
BOOM = os.path.join(DIR, '..', '..', 'app', 'src', '__fixtures__', 'boom-live')

SAMPLES = [('Balance Sheet', 'D1'), ('Balance Sheet', 'E3'), ('Balance Sheet', 'G3'),
           ('Balance Sheet', 'A5'), ('Balance Sheet', 'E6')]

def describe(c):
    f, b, fl, al = c.font, c.border, c.fill, c.alignment
    return dict(font=f'{f.name} {f.sz} bold={bool(f.b)} italic={bool(f.i)} colour={f.color.rgb if f.color else None}',
                fill=f'{fl.fill_type}/{fl.fgColor.rgb if fl.fgColor else None}',
                border=f'top={b.top.style} bottom={b.bottom.style} left={b.left.style} right={b.right.style}',
                number_format=c.number_format,
                alignment=f'horizontal={al.horizontal} indent={al.indent}')

def style_key(c):
    return (describe(c)['font'], describe(c)['fill'], describe(c)['border'],
            c.number_format, describe(c)['alignment'])

def main():
    facts = json.load(open(os.path.join(DIR, 'facts.json')))
    src = openpyxl.load_workbook(SRC)
    src_styles = set()
    for ws in src.worksheets:
        for row in ws.iter_rows():
            for c in row:
                src_styles.add(style_key(c))
    fails, printed = [], False
    for rel in facts['relationships']:
        p = os.path.join(OUT, rel['folder'],
                         f"{rel['legalName']} - Company Prepared Financial Statements FY2025.xlsx")
        wb = openpyxl.load_workbook(p)
        assert wb.sheetnames == src.sheetnames, f"{rel['folder']}: sheet names changed"
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if style_key(c) not in src_styles:
                        fails.append(f"{rel['folder']} {ws.title}!{c.coordinate}: style not present in the source")
        # arithmetic: the change column
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for r in range(4, ws.max_row + 1):
                e, g, i = ws[f'E{r}'].value, ws[f'G{r}'].value, ws[f'I{r}'].value
                if isinstance(e, (int, float)) and isinstance(g, (int, float)):
                    if abs(round(e - g, 2) - round(i, 2)) > 0.005:
                        fails.append(f"{rel['folder']} {sheet} row {r}: change column off")
        if not printed:
            print(f"Formatting comparison, source workbook against {rel['folder']} output "
                  f"(same coordinate, same role):")
            for sheet, coord in SAMPLES:
                a, b = describe(src[sheet][coord]), describe(wb[sheet][coord])
                same = a == b
                print(f"  {sheet}!{coord}  {'MATCH' if same else 'DIFFERS'}")
                for k in ('font', 'fill', 'border', 'number_format', 'alignment'):
                    print(f"      {k:14s} source: {a[k]}")
                    if a[k] != b[k]:
                        print(f"      {'':14s} output: {b[k]}")
                if not same: fails.append(f'sample {sheet}!{coord} differs')
            printed = True
    print()

    # ---- Piedmont against the Boom fixture, line by line
    spread = json.load(open(os.path.join(BOOM, 'spread-piedmont.json')))['spread']
    ratios = json.load(open(os.path.join(BOOM, 'ratios-piedmont.json')))
    pied = next(r for r in facts['relationships'] if r['folder'] == 'piedmont')
    got = {}
    for grp in ('income', 'balance', 'cash'):
        for row in pied['statements'][grp]:
            if 'v25' in row: got[row['label'].strip().lower()] = (row['v25'], row['v24'])
    diffs, checked = [], 0
    for fs in spread['financialStatements']:
        pm = {p['id']: p['endDate'] for p in fs['periods']}
        p25 = [k for k, v in pm.items() if v == '2025-12-31'][0]
        p24 = [k for k, v in pm.items() if v == '2024-12-31'][0]
        for it in fs['lineItems']:
            if it['hierarchy'] == 'header': continue
            key = it['name'].strip().lower()
            want = (it['periodValues'][p25], it['periodValues'][p24])
            if key not in got:
                diffs.append(f"MISSING  {fs['statementType']:22s} {it['name']}"); continue
            g = got[key]
            # Boom flips the sign of a line it stores as a positive deduction
            gv = (-g[0], -g[1]) if it['flipSign'] else g
            checked += 1
            if (round(gv[0]), round(gv[1])) != (round(want[0]), round(want[1])):
                diffs.append(f"DIFF     {fs['statementType']:22s} {it['name']}: "
                             f"boom {want} vs statements {gv}")
    print(f"Piedmont against the Boom fixture: {checked} line items compared, {len(diffs)} differences")
    for d in diffs: print('  ' + d)
    if diffs: fails.append('piedmont differs from the Boom fixture')
    ebitda = pied['totals']['ebitda'][0]
    for label, a, b in [('revenue', pied['totals']['revenue'][0], ratios['raw']['revenue']),
                        ('revenue prior', pied['totals']['revenue'][1], ratios['raw']['revenuePrior']),
                        ('EBITDA', ebitda, ratios['raw']['ebitda']),
                        ('total debt', pied['totals']['totalDebt'][0], ratios['raw']['totalDebt'])]:
        ok = round(a) == round(b)
        print(f"  ratios fixture {label:14s} {a:>12,} vs {b:>12,}  {'MATCH' if ok else 'DIFFERS'}")
        if not ok: fails.append(f'piedmont {label} differs from the ratios fixture')

    # ---- em dash sweep
    bad = []
    for root, _, files in os.walk(OUT):
        for fn in files:
            if fn.endswith(('.md', '.html')):
                if '—' in open(os.path.join(root, fn), encoding='utf8').read():
                    bad.append(os.path.join(root, fn))
    for rel in facts['relationships']:
        p = os.path.join(OUT, rel['folder'],
                         f"{rel['legalName']} - Company Prepared Financial Statements FY2025.xlsx")
        wb = openpyxl.load_workbook(p)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and '—' in c.value: bad.append(f'{p}!{c.coordinate}')
    print(f"\nEm dash sweep over the markdown, html and workbook text: {len(bad)} hits")
    if bad: fails += bad

    if fails:
        print('\nVERIFY FAILED'); [print('  ' + str(f)) for f in fails[:40]]; sys.exit(1)
    print('\nVERIFY PASSED')

main()
